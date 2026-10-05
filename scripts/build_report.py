"""Create the Markdown and python-docx report from executed repository evidence.

Run: python scripts/build_report.py
Does not fit models, alter inputs, or rerun the research experiments.
"""

from datetime import date
import json
from pathlib import Path
import re
import sys
from zipfile import ZipFile

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src.production import check_immutable_inputs, rate_summary, validate_december, validate_submission

REPORT_DIR = ROOT / "report"
DOCX_PATH = REPORT_DIR / "Freight_Rate_ML_Assessment_Report.docx"


def executed_evidence():
    check_immutable_inputs()
    train = pd.read_csv(ROOT / "data/train_test.csv")
    validation = pd.read_csv(ROOT / "data/validation.csv")
    results = pd.read_csv(ROOT / "reports/model_comparison.csv")
    metadata = json.loads((ROOT / "artifacts/training_metadata.json").read_text())
    preprocessing = json.loads((ROOT / "artifacts/preprocessing.json").read_text())
    summary = json.loads((ROOT / "artifacts/prediction_summary.json").read_text())
    config = json.loads((ROOT / "docs/final_model_config.json").read_text())
    eda = json.loads((ROOT / "reports/eda_summary.json").read_text())
    predictions = pd.read_csv(ROOT / "validation_predictions.csv")
    template = pd.read_csv(ROOT / "data/validation_predictions_template.csv")
    december = pd.read_csv(ROOT / "data/december_chart_inputs.csv")
    validate_submission(predictions, template)
    validate_december(december, original_inputs=preprocessing["december_original_inputs"])
    for role in ("primary", "december"):
        spec = config[role]
        rows = results[results.model.eq(spec["confirmation_model"])].sort_values("fold")
        if rows.fold.tolist() != ["fold_1", "fold_2", "fold_3"]:
            raise ValueError("Report requires all three executed fixed-budget folds")
        for metric in ("MAE", "RMSE", "R2"):
            assert np.isclose(rows[metric].mean(), spec["local_metrics_mean"][metric], rtol=1e-10)
            assert np.isclose(rows.iloc[-1][metric], spec["local_metrics_primary_holdout"][metric], rtol=1e-10)
        saved = metadata if role == "primary" else metadata["december"]
        assert saved["training_rows"] == 48000 and saved["selected_features"] == spec["feature_columns"]
        assert saved["parameters"] == spec["parameters"]
    for name, output in (("validation", predictions), ("december", december)):
        for key, value in rate_summary(output.predicted_rate).items():
            assert np.isclose(value, summary[name][key], rtol=1e-12)
    tests = re.findall(r"\*\*(\d+) tests passed, (\d+) subtests passed\*\*", (ROOT / "docs/worklog.md").read_text())
    if not tests:
        raise ValueError("Missing actual test execution evidence in the worklog")
    return train, validation, results, metadata, summary, config, eda, tests[-1]


def markdown_report(evidence):
    train, validation, results, metadata, summary, config, eda, tests = evidence
    primary, chart = config["primary"], config["december"]
    mean, holdout = primary["local_metrics_mean"], primary["local_metrics_primary_holdout"]
    city_names = set(train.pickup) | set(train.delivery)
    new_city_rows = (~validation.pickup.isin(city_names) | ~validation.delivery.isin(city_names)).sum()
    train_routes = train.pickup + " -> " + train.delivery
    final_routes = validation.pickup + " -> " + validation.delivery
    unseen_route_rows = (~final_routes.isin(train_routes)).sum()
    new_cities = sorted((set(validation.pickup) | set(validation.delivery)) - city_names)
    assert len(train) == 48000 and len(validation) == 12000
    comparisons = []
    for label, name in (
        ("Global median", "global_median"), ("Distance linear", "distance_linear"), ("Ridge", "ridge"),
        ("F-clean/direct, 284 trees", "catboost_fixed_F_clean_direct"),
        ("Selected F-clean/log, 399 trees", primary["confirmation_model"]),
        ("December C-clean/direct, 541 trees", chart["confirmation_model"]),
    ):
        averages = results[results.model.eq(name)][["MAE", "RMSE", "R2"]].mean()
        comparisons.append(f"| {label} | {averages.MAE:,.2f} | {averages.RMSE:,.2f} | {averages.R2:.4f} |")
    folds = metadata["chronological_validation_metrics"]["folds"]
    fold_rows = "\n".join(
        f"| {i} | Jan 1–{pd.Timestamp(f['train_end']).strftime('%b %d')} | "
        f"{pd.Timestamp(f['validation_start']).strftime('%b %d')}–{pd.Timestamp(f['validation_end']).strftime('%b %d')} | "
        f"{f['train_rows']:,} / {f['validation_rows']:,} |" for i, f in enumerate(folds, 1)
    )
    fold_metrics = "\n".join(
        f"| {pd.Timestamp(f['validation_start']).strftime('%b')}–{pd.Timestamp(f['validation_end']).strftime('%b')} | "
        f"{f['MAE']:.2f} | {f['RMSE']:.2f} | {f['R2']:.4f} |" for f in folds
    )
    text = f"""# Freight Rate ML Assessment

Technical report · Executed repository results · {date.today():%d %B %Y}

## 1. Executive Summary

A CatBoost solution uses **{len(train):,} labeled loads**. The frozen 399-tree log-target model achieved mean chronological **MAE ${mean['MAE']:.2f}, RMSE ${mean['RMSE']:.2f}, R² {mean['R2']:.4f}**; Sep–Oct MAE was ${holdout['MAE']:.2f}. Production fitting then used all labeled rows and generated 12,000 final predictions.

The original scorer accepted both outputs and created the embedded December chart. Local metrics describe historical validation; final accuracy and Spotter's hidden metrics are unavailable. A separate validated December model uses only available inputs.

## 2. Dataset Overview

| File | Rows × columns | Dates / role |
| --- | --- | --- |
| train_test.csv | 48,000 × 14 | Jan 1–Oct 31, 2025; labeled development |
| validation.csv | 12,000 × 13 | Nov 1–Dec 31, 2025; final inference |
| validation_predictions_template.csv | 12,000 × 2 | Unique load IDs; original rate cells remain blank |
| december_chart_inputs.csv | 31 × 7 | Every day of December 2025; fixed scenario |

The target is `posted_rate` in dollars; predictor schemas otherwise match. `load_id` is excluded from features. Raw datasets and the supplied scorer are protected by recorded SHA-256 hashes.

## 3. Data Quality Issues

| Issue | Development rows | Final-inference rows |
| --- | --- | --- |
| Missing weight | {train.weight.isna().sum():,} | {validation.weight.isna().sum():,} |
| Negative weight | {train.weight.lt(0).sum():,} | {validation.weight.lt(0).sum():,} |
| Missing market_index | {train.market_index.isna().sum():,} | {validation.market_index.isna().sum():,} |
| Missing quote_signal | {train.quote_signal.isna().sum():,} | {validation.quote_signal.isna().sum():,} |

No row was deleted. Negative weights may reflect sign corruption; their cause is unverified. Features use `abs(weight)` plus original missing/negative flags. CatBoost handles numeric NaNs natively; missing categories become explicit strings. Malformed nonmissing values and infinities fail explicitly. No duplicate load IDs or exact duplicate loads were found.

Final inference has **{len(new_cities)} new cities**, affecting **{new_city_rows:,} rows ({new_city_rows / len(validation):.2%})**, and **{eda['unseen_routes']:,} new routes**, affecting **{unseen_route_rows:,} rows ({unseen_route_rows / len(validation):.2%})**. Coordinates are retained; CatBoost accepts new labels and sklearn uses `handle_unknown="ignore"`. Compatibility does not establish future-city accuracy.

## 4. Exploratory Findings

Distance strongly tracks total rate (Pearson correlation {eda['distance_target_pearson']:.3f}). Rates are right-tailed: minimum ${eda['target']['min']:.2f}, median ${eda['target']['50%']:,.2f}, mean ${eda['target']['mean']:,.2f}, maximum ${eda['target']['max']:,.2f}. Large observations were retained rather than automatically clipped or removed.

Median dollars per mile: Dry Van ${eda['equipment']['Dry Van']['median_rpm']:.2f}, Flatbed ${eda['equipment']['Flatbed']['median_rpm']:.2f}, Reefer ${eda['equipment']['Reefer']['median_rpm']:.2f}. Calendar/load mix also varies. These associations are not causal premiums; market/quote usefulness requires chronological experiments.

<!-- page-break -->

## 5. Validation Strategy

Three expanding chronological folds reproduce the task's ordering: learn from earlier loads, then predict the next two months. All dates below are in **2025**. Every row of a calendar date remains together; training never crosses into its later validation period. Final Nov–Dec inputs are excluded from model selection, tuning, and accuracy metrics.

| Fold | Training period | Validation period | Train / validation rows |
| --- | --- | --- | --- |
{fold_rows}

Fold 3 is the primary recent holdout: training **2025-01-01–2025-08-31**, validation **2025-09-01–2025-10-31**. A random split would mix future observations and calendar/market regimes into training. MAE is primary; RMSE and R² report large-error behavior and fit on the original dollar scale.

Actual primary-model fixed-budget results:

| Validation | MAE ($) | RMSE ($) | R² |
| --- | --- | --- | --- |
{fold_metrics}
| Unweighted fold mean | {mean['MAE']:.2f} | {mean['RMSE']:.2f} | {mean['R2']:.4f} |

All learned sklearn imputation, scaling, and categorical vocabularies are fitted on each applicable training fold. CatBoost's learned categorical statistics are training-fold local; missing numeric handling is native. Feature generation is stateless and row-wise.

Initial experiments used the outer chronological holdout for early stopping. Shortlisted candidates were confirmed with fixed median tree budgets, without an eval_set. Historical folds were reused for selection, so these are selection estimates, not an untouched test. Fold averages are unweighted and are not pooled metrics.

## 6. Feature Engineering

One shared builder in `src/features.py` serves both training and inference. The primary ordered whitelist has **26 features**; it excludes `load_id`, `posted_rate`, and `predicted_rate`. No external target encoding is used. The calendar reference remains 2025-01-01.

| Group | Selected feature names |
| --- | --- |
| Operational | distance, equipment, weight_clean, weight_missing, weight_negative |
| Geographic | pickup, delivery, pickup_lat, pickup_lon, delivery_lat, delivery_lon |
| Route | route = pickup + " -> " + delivery |
| Calendar | month, day, day_of_week, day_of_year, iso_week, weekend, days_since_reference_date |
| Cyclical | sin_day_of_week, cos_day_of_week, sin_day_of_year, cos_day_of_year |
| Market | market_index, market_index_missing, quote_signal |

`weight_clean` is absolute original weight; missing cleaned weights remain NaN. Missing/negative flags describe the original measurement. City/route strings support native categories, while coordinates supply information when labels are unseen. The December model retains operational, geographic, calendar, and cyclic features: **22 features**, excluding route and market/quote fields.

Evidence: `docs/validation_plan.md`, `docs/final_model_config.json`, `reports/model_comparison.csv`, and `src/features.py`.

<!-- page-break -->

## 7. Model Experiments

Executed comparisons include global median, distance-only linear regression, Ridge, and three modest CatBoost settings (depths 6/7/8). Controlled experiments add operational/calendar → cities → coordinates → route → market_index → quote_signal on the same folds, and compare raw weight with absolute weight plus flags, and direct with log1p targets. No large search was performed.

Mean chronological metrics for the baselines and fixed-budget finalists:

| Model | MAE ($) | RMSE ($) | R² |
| --- | --- | --- | --- |
{chr(10).join(comparisons)}

Market features improve aggregate MAE, but their incremental benefit varies by fold. Quote gains are modest and mixed. The full selected bundle won among executed configurations; route's contribution conditional on that final bundle was not isolated. Source provenance and prediction-time availability of market/quote signals remain limitations, not demonstrated leakage.

## 8. Final Model Selection

The primary is **CatBoostRegressor, 399 trees, depth 6, learning_rate 0.05, l2_leaf_reg 5, MAE loss, random_seed 42, thread_count 4, nan_mode Min**. It fits `log1p(posted_rate)` and converts predictions with `expm1`. The feature list and tree budget were frozen before full-data production fitting.

Selection prioritized fixed-budget mean chronological MAE, recent holdout behavior, stability, missing/unseen robustness, and realistic inputs. The selected model beats every simple baseline. Log modeling improves mean MAE from $131.00 to $129.41 versus the fixed direct-target alternative, but slightly worsens recent MAE and mean RMSE; it does not win every metric. Its Sep–Oct results are **MAE ${holdout['MAE']:.2f}, RMSE ${holdout['RMSE']:.2f}, R² {holdout['R2']:.4f}**.

## 9. Final Prediction Process

Both frozen models train on all 48,000 development rows with no eval_set or early stopping. Native models, preprocessing specifications, deterministic city coordinates, feature/source hashes, parameters, versions, and historical metrics are saved under `artifacts/`.

The saved primary predicts every final-inference load. Predictions join to template IDs **one-to-one by load_id**, in template order, with exactly `load_id,predicted_rate` and no index. Checks reject missing/duplicate/extra IDs and nonfinite or nonpositive rates. The frozen $0.01 guard altered zero production rates.

| Output | Rows | Minimum ($) | Maximum ($) | Mean ($) | Median ($) |
| --- | --- | --- | --- | --- | --- |
| Final inference | {summary['validation']['rows']:,} | {summary['validation']['minimum']:,.2f} | {summary['validation']['maximum']:,.2f} | {summary['validation']['mean']:,.2f} | {summary['validation']['median']:,.2f} |
| December | {summary['december']['rows']} | {summary['december']['minimum']:,.2f} | {summary['december']['maximum']:,.2f} | {summary['december']['mean']:,.2f} | {summary['december']['median']:,.2f} |

Native-model round trips reproduce predictions exactly. The executed production rerun reproduced both CSV hashes exactly. Output summaries are not accuracy metrics.

<!-- page-break -->

## 10. December Scenario

The official chart fixes **Lexington → Fort Wayne, 360 miles, Dry Van, 32,000 lb**, while the date changes from **2025-12-01 through 2025-12-31**. All original six inputs and the seven-column CSV order are preserved; only `predicted_rate` is filled.

![Official scorer chart: fixed freight characteristics, changing December date](../scorer_results/candidate_december.png)

The separate frozen **541-tree group-C direct-target model** uses the same chronological philosophy and only available fields. Coordinates come from consistent city lookups in supplied development data; no future market_index or quote_signal is invented. Its mean historical MAE/RMSE/R² are **${chart['local_metrics_mean']['MAE']:.2f} / ${chart['local_metrics_mean']['RMSE']:.2f} / {chart['local_metrics_mean']['R2']:.4f}**. These do not establish accuracy for the fixed December lane. Actual December rates range from **${summary['december']['minimum']:.2f} to ${summary['december']['maximum']:.2f}**.

The original `score.py` validated 12,000 final predictions and 31 fixed December predictions and created this chart. It validates contracts; hidden accuracy is calculated by Spotter after submission.

## 11. Limitations and Improvements

- Reused historical folds can overstate generalization; an independent later labeled period or nested chronological stopping split would strengthen evaluation.
- November/December labels and unseen-city accuracy are unavailable. Ten months do not establish full-year seasonality; trees have limited extrapolation beyond observed calendar ranges.
- Large-rate errors remain: RMSE is substantially above MAE. Investigate tail provenance and performance by lane/equipment before changing target treatment.
- Verify market/quote timestamps and availability, supplied coordinate accuracy, and weight-error origins. Their current uncertainty is documented rather than assumed away.
- Future work could monitor temporal/input drift and calibration. These are proposed improvements, not completed experiments.

## 12. Reproducibility

Use Python {metadata['package_versions']['python']} and the pinned tested versions in `constraints.txt`, alongside compatible requirements. From the repository root:

```bash
python -m venv .venv
source .venv/bin/activate
python -m pip install -r requirements.txt -c constraints.txt
python run_pipeline.py
pytest -q
python score.py --predictions validation_predictions.csv \\
  --december-predictions data/december_chart_inputs.csv
python scripts/build_report.py
```

The pipeline validates, trains, saves, predicts, and checks integrity without rerunning research. Recorded verification: **{tests[0]} tests plus {tests[1]} subtests passed**, pipeline passed, original scorer passed. Protected inputs remain unchanged. `README.md`, `docs/worklog.md`, and metadata provide the audit trail. Publication and an actual Loom recording are separate submission steps.
"""
    return text


def shade(cell, fill):
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(element)


def add_inline(paragraph, value):
    for piece in re.split(r"(\*\*.*?\*\*|`[^`]+`)", value):
        if not piece:
            continue
        if piece.startswith("**"):
            paragraph.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`"):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = "DejaVu Sans Mono"
            run.font.size = Pt(9)
        else:
            paragraph.add_run(piece)


def add_table(document, lines):
    rows = [[item.strip() for item in line.strip().strip("|").split("|")] for line in lines]
    rows = [row for row in rows if not all(re.fullmatch(r":?-+:?", item) for item in row)]
    table = document.add_table(rows=0, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = {2: [1.2, 5.8], 3: [2.7, 1.25, 3.05], 4: [3.2, 1.27, 1.27, 1.26], 6: [1.3, .65, 1.3, 1.3, 1.3, 1.15]}
    if rows[0][0] == "Fold":
        widths[4] = [.60, 1.90, 1.90, 2.60]
    elif rows[0][0] in ("Validation",):
        widths[4] = [2.8, 1.4, 1.4, 1.4]
    elif rows[0][0] == "Issue":
        widths[3] = [3.1, 1.95, 1.95]
    for column, width in zip(table.columns, widths[len(rows[0])]):
        column.width = Inches(width)
    for index, values in enumerate(rows):
        row = table.add_row()
        properties = row._tr.get_or_add_trPr()
        properties.append(OxmlElement("w:cantSplit"))
        if index == 0:
            properties.append(OxmlElement("w:tblHeader"))
        for cell, value, width in zip(row.cells, values, widths[len(values)]):
            cell.width = Inches(width)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.space_before = Pt(3)
            add_inline(paragraph, value)
            for run in paragraph.runs:
                run.font.size = Pt(9.1)
                if index == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string("FFFFFF")
            if index == 0:
                shade(cell, "17384A")
            elif index % 2:
                shade(cell, "F1F5F7")
    after = document.add_paragraph()
    after.paragraph_format.space_after = Pt(0)
    after.paragraph_format.space_before = Pt(0)
    after.paragraph_format.line_spacing = Pt(3)
    after.add_run().font.size = Pt(1)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), instruction)
    run._r.addnext(field)


def create_docx(markdown):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Inches(.65), Inches(.65)
    section.left_margin, section.right_margin = Inches(.75), Inches(.75)
    section.header_distance, section.footer_distance = Inches(.25), Inches(.25)
    normal = document.styles["Normal"]
    normal.font.name = "DejaVu Sans"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string("24313A")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    heading = document.styles["Heading 1"]
    heading.font.name, heading.font.size = "DejaVu Sans", Pt(12)
    heading.font.color.rgb = RGBColor.from_string("17384A")
    heading.paragraph_format.space_before = Pt(9)
    heading.paragraph_format.space_after = Pt(4)
    heading.paragraph_format.keep_with_next = True
    title = document.styles["Title"]
    title.font.name, title.font.size = "DejaVu Sans", Pt(22)
    title.font.color.rgb = RGBColor.from_string("17384A")
    title.paragraph_format.space_after = Pt(4)
    header = section.header.paragraphs[0]
    header.add_run("FREIGHT RATE ML ASSESSMENT  |  TECHNICAL REPORT").font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Executed repository results  •  Page ").font.size = Pt(8)
    add_field(footer, "PAGE")
    footer.add_run(" of ").font.size = Pt(8)
    add_field(footer, "NUMPAGES")
    document.core_properties.title = "Freight Rate ML Assessment Report"
    document.core_properties.subject = "Chronological freight-rate validation and fixed December scenario"
    document.core_properties.author = ""
    lines = markdown.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        index += 1
        if not line:
            continue
        if line == "<!-- page-break -->":
            document.add_page_break()
        elif line.startswith("# "):
            document.add_paragraph(line[2:], style="Title")
        elif line.startswith("## "):
            document.add_paragraph(line[3:], style="Heading 1")
        elif line.startswith("Technical report"):
            paragraph = document.add_paragraph(line)
            paragraph.runs[0].italic = True
            paragraph.runs[0].font.size = Pt(9)
        elif line.startswith("|"):
            table_lines = [line]
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            add_table(document, table_lines)
        elif line.startswith("```"):
            paragraph = document.add_paragraph()
            while index < len(lines) and not lines[index].startswith("```"):
                run = paragraph.add_run(lines[index] + "\n")
                run.font.name, run.font.size = "DejaVu Sans Mono", Pt(8.2)
                index += 1
            index += 1
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.line_spacing = 1
        elif line.startswith("!["):
            match = re.fullmatch(r"!\[(.*?)\]\((.*?)\)", line)
            paragraph = document.add_paragraph()
            picture = paragraph.add_run().add_picture(str((REPORT_DIR / match.group(2)).resolve()), width=Inches(7))
            picture._inline.docPr.set("descr", match.group(1))
            paragraph.paragraph_format.space_after = Pt(4)
        elif line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, line[2:])
        else:
            paragraph = document.add_paragraph()
            add_inline(paragraph, line)
    document.save(DOCX_PATH)


def verify_docx():
    reopened = Document(DOCX_PATH)
    headings = [paragraph.text for paragraph in reopened.paragraphs if paragraph.style.name == "Heading 1"]
    if len(headings) != 12 or not all(heading.startswith(f"{number}. ") for number, heading in enumerate(headings, 1)):
        raise ValueError("Report is missing a required numbered section")
    if len(reopened.inline_shapes) != 1:
        raise ValueError("Report must embed the official December chart")
    with ZipFile(DOCX_PATH) as archive:
        if archive.testzip() is not None:
            raise ValueError("DOCX ZIP archive is corrupt")
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
        original = (ROOT / "scorer_results/candidate_december.png").read_bytes()
        if not any(archive.read(name) == original for name in media):
            raise ValueError("Embedded chart differs from the supplied scorer output")
    print(f"DOCX reopened successfully: {len(headings)} sections, {len(reopened.tables)} tables, original chart embedded.")
    print(f"Saved {REPORT_DIR / 'report.md'} and {DOCX_PATH}")


def main():
    REPORT_DIR.mkdir(exist_ok=True)
    markdown = markdown_report(executed_evidence())
    (REPORT_DIR / "report.md").write_text(markdown, encoding="utf-8")
    create_docx(markdown)
    verify_docx()


if __name__ == "__main__":
    main()
